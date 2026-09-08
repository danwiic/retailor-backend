import io
import os
from datetime import datetime, timezone

import boto3
from botocore.config import Config
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BUCKET = os.getenv("S3_BUCKET", "resume-tailor-exports")
EXPORT_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"

FONT_NAME = "Arial"
BODY_SIZE = Pt(10.5)
SECTION_SIZE = Pt(10.5)
NAME_SIZE = Pt(18)
TEXT_COLOR = RGBColor(0x00, 0x00, 0x00)
MUTED_COLOR = RGBColor(0x50, 0x50, 0x50)
HEADER_COLOR = "1F4E79"
RULE_COLOR = "BFBFBF"

# Page geometry — keep these in sync so tab stops and table widths always
# reach the true right margin instead of stopping short.
MARGIN = Inches(0.55)
PAGE_WIDTH = Inches(8.5)
CONTENT_WIDTH = PAGE_WIDTH - (2 * MARGIN)  # 7.4"

SECTION_SPACE_BEFORE = 9
SECTION_SPACE_AFTER = 3
ENTRY_SPACE_BEFORE = 5


def _set_run_font(run, size=BODY_SIZE, bold=False, italic=False, color=TEXT_COLOR):
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = size
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _set_char_spacing(run, twips):
    """Add letter-spacing (tracking) to a run — gives headers a more
    deliberate, designed feel instead of default cramped kerning."""
    rpr = run._element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twips))
    rpr.append(spacing)


def _set_compact_paragraph(paragraph, before=0, after=0, line=1.08):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.line_spacing = line


def _add_border(paragraph, position="bottom", size=6, color=HEADER_COLOR, space=1):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    edge = OxmlElement(f"w:{position}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)
    borders.append(edge)


def _add_right_tab(paragraph):
    tabs = paragraph.paragraph_format.tab_stops
    tabs.add_tab_stop(CONTENT_WIDTH, alignment=2)


def _add_section(doc, title):
    paragraph = doc.add_paragraph()
    _set_compact_paragraph(
        paragraph, before=SECTION_SPACE_BEFORE, after=SECTION_SPACE_AFTER
    )
    _add_border(paragraph, size=6, color=HEADER_COLOR)
    run = paragraph.add_run(title.upper())
    _set_run_font(
        run, size=SECTION_SIZE, bold=True, color=RGBColor.from_string(HEADER_COLOR)
    )
    _set_char_spacing(run, 12)
    return paragraph


def _add_metadata_line(
    doc, left_parts, right_text="", bold_first=False, italic_right=False, space_before=0
):
    paragraph = doc.add_paragraph()
    _set_compact_paragraph(paragraph, before=space_before, after=0)
    _add_right_tab(paragraph)
    left_parts = [part for part in left_parts if part]
    for index, part in enumerate(left_parts):
        if index:
            separator = paragraph.add_run("  •  ")
            _set_run_font(separator, color=MUTED_COLOR)
        run = paragraph.add_run(part)
        _set_run_font(run, bold=bold_first and index == 0)
    if right_text:
        run = paragraph.add_run(f"\t{right_text}")
        _set_run_font(run, italic=italic_right, color=MUTED_COLOR)
    return paragraph


def _add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    _set_compact_paragraph(paragraph, after=1)
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.first_line_indent = Inches(-0.14)
    run = paragraph.add_run(text)
    _set_run_font(run)


def build_docx(resume: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = MARGIN
    section.right_margin = MARGIN

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    name = doc.add_paragraph()
    _set_compact_paragraph(name, after=2)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name.add_run(resume.get("name", "Resume"))
    _set_run_font(name_run, size=NAME_SIZE, bold=True)
    _set_char_spacing(name_run, 16)

    contact = resume.get("contact") or []
    contact_line = "   |   ".join(c.get("value", "") for c in contact if c.get("value"))
    if contact_line:
        contact_paragraph = doc.add_paragraph()
        _set_compact_paragraph(contact_paragraph, after=8)
        _add_border(contact_paragraph, size=4, color=RULE_COLOR, space=6)
        contact_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = contact_paragraph.add_run(contact_line)
        _set_run_font(run, color=MUTED_COLOR)
    else:
        # Still separate the header block from the body even with no contact info.
        spacer = doc.add_paragraph()
        _set_compact_paragraph(spacer, after=8)
        _add_border(spacer, size=4, color=RULE_COLOR, space=6)

    summary = resume.get("summary")
    if summary:
        _add_section(doc, "Summary")
        paragraph = doc.add_paragraph()
        _set_compact_paragraph(paragraph)
        _set_run_font(paragraph.add_run(summary))

    skills = resume.get("skills")
    if skills:
        _add_section(doc, "Technical Skills")
        paragraph = doc.add_paragraph()
        _set_compact_paragraph(paragraph)
        _set_run_font(paragraph.add_run(" • ".join(skills)))

    def add_entries(heading, items):
        if not items:
            return
        _add_section(doc, heading)
        for i, item in enumerate(items):
            title = item.get("title") or item.get("name", "")
            company = item.get("company", "")
            dates = item.get("dates", "")
            _add_metadata_line(
                doc,
                [title, company],
                dates,
                bold_first=True,
                italic_right=True,
                space_before=ENTRY_SPACE_BEFORE if i > 0 else 0,
            )
            for bullet in item.get("bullets") or []:
                _add_bullet(doc, bullet)

    add_entries("Experience", resume.get("experience"))
    add_entries("Projects", resume.get("projects"))

    education = resume.get("education")
    if education:
        _add_section(doc, "Education")
        for i, item in enumerate(education):
            school = item.get("school", "")
            degree = item.get("degree", "")
            location = item.get("location", "")
            dates = item.get("dates", "")
            _add_metadata_line(
                doc,
                [school],
                location,
                bold_first=True,
                italic_right=True,
                space_before=ENTRY_SPACE_BEFORE if i > 0 else 0,
            )
            _add_metadata_line(doc, [degree], dates, italic_right=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pdf(resume: dict) -> bytes:
    """Build a PDF using the same compact layout as the DOCX."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    content_width = 7.4 * inch
    styles = getSampleStyleSheet()
    # ReportLab's built-in fonts do not include Arial. Register a bundled Arial
    # TTF in deployment if exact PDF font identity is required; Helvetica is
    # the metrically compatible fallback for local generation.
    body = ParagraphStyle(
        "ResumeBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=13.5,
        spaceAfter=2,
    )
    section_style = ParagraphStyle(
        "ResumeSection",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=10.5,
        textColor=colors.HexColor("#1F4E79"),
        spaceBefore=0,
        spaceAfter=0,
        leading=13,
    )
    muted = colors.HexColor("#505050")
    small_center = ParagraphStyle(
        "ResumeContact",
        parent=body,
        alignment=TA_CENTER,
        textColor=muted,
        spaceAfter=0,
        leading=13,
    )
    title_style = ParagraphStyle(
        "ResumeName",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=20,
        alignment=TA_CENTER,
        spaceAfter=4,
    )
    right_style = ParagraphStyle(
        "Right",
        parent=body,
        alignment=TA_RIGHT,
        fontName="Helvetica-Oblique",
        textColor=muted,
    )

    story = [Paragraph(resume.get("name", "Resume"), title_style)]
    contact = resume.get("contact") or []
    contact_line = "   |   ".join(c.get("value", "") for c in contact if c.get("value"))
    if contact_line:
        story.append(Paragraph(contact_line, small_center))
    story.append(Spacer(1, 6))
    rule = Table([[""]], colWidths=[content_width], rowHeights=[1])
    rule.setStyle(
        TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.75, colors.HexColor("#BFBFBF"))])
    )
    story.append(rule)
    story.append(Spacer(1, 8))

    def section(title):
        table = Table(
            [[Paragraph(title.upper(), section_style)]], colWidths=[content_width]
        )
        table.setStyle(
            TableStyle(
                [
                    ("LINEBELOW", (0, 0), (-1, -1), 0.75, colors.HexColor("#1F4E79")),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 3))

    def metadata(left, right, bold=True, space_before=0):
        if space_before:
            story.append(Spacer(1, space_before))
        left_markup = f"<b>{left}</b>" if bold else left
        table = Table(
            [
                [
                    Paragraph(left_markup, body),
                    Paragraph(right.replace("\n", "<br/>"), right_style),
                ]
            ],
            colWidths=[content_width - 2.3 * inch, 2.3 * inch],
        )
        table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ]
            )
        )
        story.append(table)

    if resume.get("summary"):
        section("Summary")
        story.append(Paragraph(resume["summary"], body))
        story.append(Spacer(1, 6))
    if resume.get("skills"):
        section("Technical Skills")
        story.append(Paragraph(" • ".join(resume["skills"]), body))
        story.append(Spacer(1, 6))

    def entries(title, items):
        if not items:
            return
        section(title)
        for i, item in enumerate(items):
            name = item.get("title") or item.get("name", "")
            metadata(
                " • ".join(x for x in [name, item.get("company", "")] if x),
                item.get("dates", ""),
                space_before=5 if i > 0 else 0,
            )
            for bullet in item.get("bullets") or []:
                story.append(Paragraph(f"&bull;&nbsp;&nbsp;{bullet}", body))
        story.append(Spacer(1, 6))

    entries("Experience", resume.get("experience"))
    entries("Projects", resume.get("projects"))
    if resume.get("education"):
        section("Education")
        for i, item in enumerate(resume["education"]):
            metadata(
                item.get("school", ""),
                item.get("location", ""),
                space_before=5 if i > 0 else 0,
            )
            metadata(item.get("degree", ""), item.get("dates", ""), bold=False)

    doc.build(story)
    return buf.getvalue()


def _s3_client():
    return boto3.client(
        "s3",
        region_name=os.getenv("AWS_REGION", "us-east-1"),
        aws_access_key_id=os.getenv("AWS_ACCESS_KEY_ID"),
        aws_secret_access_key=os.getenv("AWS_SECRET_ACCESS_KEY"),
        config=Config(connect_timeout=10, read_timeout=30, retries={"max_attempts": 2}),
    )


def upload_export_and_sign(
    content: bytes, key: str, content_type: str = EXPORT_MIME
) -> str:
    s3 = _s3_client()
    s3.put_object(Bucket=BUCKET, Key=key, Body=content, ContentType=content_type)
    return s3.generate_presigned_url(
        "get_object",
        Params={"Bucket": BUCKET, "Key": key},
        ExpiresIn=int(os.getenv("SIGNED_URL_TTL", "300")),
    )


def export_key(job_id: str, export_format: str = "docx") -> str:
    ts = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    return f"{job_id}/{ts}_tailored.{export_format}"
